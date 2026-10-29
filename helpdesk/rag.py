"""RAG pipeline: load documents from docs/raw, chunk, embed with Ollama, store in ChromaDB, hybrid search.

    python -m helpdesk.rag index            # (re)build the vector store from docs/raw
    python -m helpdesk.rag search "query"   # inspect what retrieval returns for a query
"""
import csv
import math
import re
import sys
from collections import Counter
from functools import lru_cache
from html.parser import HTMLParser
from pathlib import Path

import chromadb
import ollama
from chromadb.config import Settings

from . import config

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".html", ".htm", ".csv"}


# ---------- loaders: each returns [(page_number or 0, text)] ----------

def _load_pdf(path):
    from pypdf import PdfReader
    pages = []
    for i, page in enumerate(PdfReader(path).pages):
        text = page.extract_text() or ""
        lines = [ln for ln in text.splitlines() if ln.strip()]
        # some PDFs (e.g. reg_2019.pdf) extract one word per line; join those back into running text
        if lines and sum(len(ln.split()) for ln in lines) / len(lines) < 2:
            text = re.sub(r"\s+", " ", text)
        pages.append((i + 1, text))
    return pages


def _load_docx(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return [(0, "\n".join(parts))]


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts, self._skip = [], 0

    def handle_starttag(self, tag, attrs):
        self._skip += tag in ("script", "style")

    def handle_endtag(self, tag):
        self._skip -= tag in ("script", "style")
        if tag in ("p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "table"):
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def _load_html(path):
    p = _HTMLText()
    p.feed(path.read_text(errors="ignore"))
    return [(0, "".join(p.parts))]


def _load_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    return [(0, "\n".join("; ".join(f"{k}: {v}" for k, v in r.items() if v) for r in rows))]


def load(path: Path):
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _load_pdf(path)
    if ext == ".docx":
        return _load_docx(path)
    if ext in (".html", ".htm"):
        return _load_html(path)
    if ext == ".csv":
        return _load_csv(path)
    return [(0, path.read_text(errors="ignore"))]


def chunk(text: str, size=config.CHUNK_CHARS, overlap=config.CHUNK_OVERLAP):
    """Split into ~size-char chunks, preferring line/sentence boundaries, with overlap."""
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            cut = max(text.rfind("\n", start + size // 2, end), text.rfind(". ", start + size // 2, end))
            if cut > start:
                end = cut + 1
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [c for c in chunks if c]


# ---------- vector store ----------

def _client():
    return chromadb.PersistentClient(path=str(config.CHROMA_DIR), settings=Settings(anonymized_telemetry=False))


def _embed(texts, prefix):
    # nomic-embed-text expects task prefixes ("search_document: " / "search_query: ")
    client = ollama.Client(host=config.OLLAMA_API_BASE)
    out = []
    for i in range(0, len(texts), 32):
        out += client.embed(model=config.EMBED_MODEL, input=[prefix + t for t in texts[i:i + 32]]).embeddings
    return out


def index():
    files = sorted(p for p in config.RAW_DIR.rglob("*") if p.is_file() and not p.name.startswith("."))
    if not files:
        sys.exit(f"No documents found in {config.RAW_DIR}")
    client = _client()
    if config.COLLECTION in [c.name for c in client.list_collections()]:
        client.delete_collection(config.COLLECTION)
    col = client.create_collection(config.COLLECTION, configuration={"hnsw": {"space": "cosine"}})

    print(f"{'file':60} {'pages':>5} {'chunks':>6}")
    total = 0
    for path in files:
        source = str(path.relative_to(config.RAW_DIR))
        if path.suffix.lower() not in SUPPORTED:
            print(f"{source:60} SKIPPED (unsupported type; convert to PDF/DOCX/TXT/CSV)")
            continue
        try:
            pages = load(path)
        except Exception as e:  # one bad file must not abort the whole index
            print(f"{source:60} FAILED: {e}")
            continue
        title = path.stem.replace("_", " ").replace("-", " ")
        ids, docs, metas = [], [], []
        for page, text in pages:
            for i, c in enumerate(chunk(text)):
                ids.append(f"{source}::p{page}::c{i}")
                docs.append(f"{title}\n{c}")
                metas.append({"source": source, "page": page})
        n_chars = sum(len(t) for _, t in pages)
        warn = ""
        if path.suffix.lower() == ".pdf" and n_chars < 50 * len(pages):
            warn = "  WARNING: little/no text extracted - scanned PDF? run OCR first (e.g. ocrmypdf)"
        if docs:
            col.add(ids=ids, documents=docs, metadatas=metas, embeddings=_embed(docs, "search_document: "))
        total += len(docs)
        print(f"{source:60} {len(pages):>5} {len(docs):>6}{warn}")
    print(f"\nIndexed {total} chunks from {len(files)} files into {config.CHROMA_DIR}")


@lru_cache(maxsize=1)
def _collection():
    try:
        return _client().get_collection(config.COLLECTION)
    except Exception:
        sys.exit("Vector store not found. Run: python -m helpdesk.rag index")


def _tokens(text):
    return re.findall(r"[a-z0-9]+", text.lower())


@lru_cache(maxsize=1)
def _bm25_index():
    # ponytail: in-memory BM25 over all chunks, fine for a few thousand chunks; use a real FTS index beyond that
    data = _collection().get(include=["documents", "metadatas"])
    toks = [_tokens(d) for d in data["documents"]]
    df = Counter(t for d in toks for t in set(d))
    avgdl = sum(map(len, toks)) / max(len(toks), 1)
    return data, toks, df, avgdl


def _bm25(query, n, k1=1.5, b=0.75):
    data, toks, df, avgdl = _bm25_index()
    N, q = len(toks), set(_tokens(query))
    scores = []
    for i, d in enumerate(toks):
        tf = Counter(d)
        s = sum(math.log(1 + (N - df[t] + 0.5) / (df[t] + 0.5)) * tf[t] * (k1 + 1)
                / (tf[t] + k1 * (1 - b + b * len(d) / avgdl)) for t in q if t in tf)
        if s > 0:
            scores.append((s, data["ids"][i]))
    return [i for _, i in sorted(scores, reverse=True)[:n]]


def search(query: str, k: int = config.TOP_K):
    """Hybrid retrieval: dense (Chroma) + BM25 (exact terms like course codes, '75%'), fused with RRF."""
    col = _collection()
    dense = col.query(query_embeddings=_embed([query], "search_query: "), n_results=min(10, col.count()))["ids"][0]
    fused = Counter()
    for ranking in (dense, _bm25(query, 10)):
        for rank, cid in enumerate(ranking):
            fused[cid] += 1 / (60 + rank)
    top = [cid for cid, _ in fused.most_common(k)]
    got = col.get(ids=top, include=["documents", "metadatas"])
    by_id = dict(zip(got["ids"], zip(got["documents"], got["metadatas"])))
    return [{"source": by_id[c][1]["source"], "page": by_id[c][1]["page"], "text": by_id[c][0]} for c in top]


if __name__ == "__main__":
    if sys.argv[1:2] == ["index"]:
        index()
    elif sys.argv[1:2] == ["search"] and len(sys.argv) > 2:
        for r in search(" ".join(sys.argv[2:])):
            print(f"--- {r['source']} (page {r['page']})\n{r['text'][:500]}\n")
    else:
        sys.exit(__doc__)
