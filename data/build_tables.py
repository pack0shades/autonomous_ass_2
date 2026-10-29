"""Build data/courses.csv, data/slots.csv, data/faculty.csv and docs/raw/timetable_slot_system_2026_27.md
from the official slot-system workbook in data/source/.

    python data/build_tables.py

Inputs (data/source/):
  slot_system_2026_27_sem1.xlsx      official AY 2026-27 Sem I slot system + classroom allocation (+ other sheets)
  faculty_emails_collected.csv       instructor emails collected from department web pages (name,department,email)
"""
import csv
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

DATA = Path(__file__).resolve().parent
SRC = DATA / "source"
DOC = DATA.parent / "docs" / "raw" / "timetable_slot_system_2026_27.md"
DAYS = {"Mon": "Monday", "Tue": "Tuesday", "Wed": "Wednesday", "Thu": "Thursday", "Fri": "Friday",
        "Sat": "Saturday", "Sun": "Sunday"}
# Table 4.3 of the UG regulations (reg_2019.pdf): third letter of a course code = course type
COURSE_TYPE = {"L": "Lecture", "P": "Practical/Lab", "D": "Project", "N": "Non-graded", "T": "Thesis",
               "Q": "Seminar", "S": "Independent Study", "R": "Engineering Innovation", "E": "Engineering Design"}
# The Time Table Committee sheet lists SoLA as "Humanities & Social Sciences / SOLA"
DEPT_RENAME = {"School of Liberal Arts (SoLA)": "School of Liberal Arts (SoLA) / Humanities and Social Sciences (HSS)"}
# The assignment's example prompts use CSL7570 for Autonomous Systems, which the official timetable lists as
# CSL7650. Kept as an explicitly marked alias row so the example prompts can be answered.
ALIASES = {"CSL7570": "CSL7650"}

wb = openpyxl.load_workbook(SRC / "slot_system_2026_27_sem1.xlsx", data_only=True)


def clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def norm_name(n):
    return " ".join(re.findall(r"[a-z]+", re.sub(r"\b(dr|prof)\b\.?", " ", n.lower())))


def people(cell):
    """'00304-Pranay Ranjan, 00597-Ravi Prakash Srivastava' -> ['Pranay Ranjan', 'Ravi Prakash Srivastava']"""
    out = []
    for p in re.split(r",|;|/|\n", clean(cell)):
        p = re.sub(r"^\d+\s*-\s*", "", p.strip()).strip(" .")
        if p and p not in out:
            out.append(p)
    return out


def hhmm(t):
    """'9– 9:50' style hour text -> ('09:00', '09:50'); afternoon hours (1-7) are PM."""
    nums = re.findall(r"(\d{1,2})(?::(\d{2}))?", t)
    (h1, m1), (h2, m2) = nums[0], nums[1]
    fix = lambda h: int(h) + 12 if int(h) < 8 else int(h)
    return f"{fix(h1):02d}:{m1 or '00'}", f"{fix(h2):02d}:{m2 or '00'}"


# ---------- slots.csv: from the "Slot System" grid ----------
grid = wb["Slot System"]
merged_rows = {}  # top-left cell -> number of rows the merge spans
for rng in grid.merged_cells.ranges:
    merged_rows[(rng.min_row, rng.min_col)] = rng.max_row - rng.min_row + 1
day_cols = {2: "Monday", 4: "Tuesday", 6: "Wednesday", 8: "Thursday", 10: "Friday"}  # B D F H J (+1 = lab column)
times = {r: hhmm(clean(grid.cell(r, 1).value)) for r in range(3, 13)}  # rows 3..12 = 8:00 .. 17:50
slots = []
for r, (start, end) in times.items():
    for col, day in day_cols.items():
        for c in (col, col + 1):
            v = clean(grid.cell(r, c).value)
            if re.fullmatch(r"[A-Z]{1,2}\d{0,2}", v):
                span = merged_rows.get((r, c), 1)
                slots.append({"slot": v, "day": day, "start_time": start, "end_time": times[r + span - 1][1]})

# M.Tech. Executive extra slots: weekday evenings O-S (6-7:30 PM), Saturday O-S, and the Y slot table
ex = wb["Extra Slots for MTech Executive"]
for i, s in enumerate("OPQRS"):
    slots.append({"slot": s, "day": list(day_cols.values())[i], "start_time": "18:00", "end_time": "19:30"})
for r in range(1, ex.max_row + 1):
    if clean(ex.cell(r, 7).value) in tuple("OPQRS") and re.match(r"\d", clean(ex.cell(r, 1).value)):
        s, e = hhmm(clean(ex.cell(r, 1).value))
        slots.append({"slot": clean(ex.cell(r, 7).value), "day": "Saturday", "start_time": s, "end_time": e})
    y = clean(ex.cell(r, 10).value)
    if re.fullmatch(r"Y[A-Z]\d?( \(for labs\))?", y):
        for day, c in (("Saturday", 11), ("Sunday", 12)):
            t = clean(ex.cell(r, c).value)
            if re.search(r"\d", t):
                s, e = re.split(r"\s*[–-]\s*", t)
                to24 = lambda x: (lambda h, m, pm: f"{int(h) % 12 + (12 if pm else 0):02d}:{m}")(
                    *re.match(r"(\d+):(\d+)", x).groups(), "PM" in x or int(x.split(":")[0]) < 8)
                slots.append({"slot": y.split()[0], "day": day, "start_time": to24(s), "end_time": to24(e)})
# ponytail: X slots (XO-XS) are skipped, their sheet layout does not say which day they run on

# ---------- courses.csv: from the "Classroom Allocation" sheet, one row per course code ----------
rows = [r[:10] for r in wb["Classroom Allocation"].iter_rows(min_row=2, values_only=True) if r[0]]
by_code = defaultdict(list)
for r in rows:
    by_code[clean(r[0]).upper()].append(r)

courses = []
for code, rs in by_code.items():
    names = sorted({clean(r[1]) for r in rs}, key=len)
    base = names[0]
    sections, rooms, instr, co = [], [], [], []
    for r in rs:
        variant = clean(r[1])[len(base):].strip(" ()") if clean(r[1]).startswith(base) else ""
        sec = clean(r[4]) + (f" ({variant})" if variant else "")
        if sec and sec not in sections:
            sections.append(sec)
        if clean(r[9]) and clean(r[9]) not in rooms:
            rooms.append(clean(r[9]))
        instr += [p for p in people(r[6]) if p not in instr]
        co += [p for p in people(r[7]) if p not in co and p not in instr]
    dept = next((clean(r[5]) for r in rs if r[5]), "")
    m = re.match(r"[A-Z]{2}([A-Z])", code)
    courses.append({
        "course_code": code, "course_name": base, "department": DEPT_RENAME.get(dept, dept),
        "course_type": COURSE_TYPE.get(m.group(1), "") if m else "",
        "ltp": clean(rs[0][2]), "credits": clean(rs[0][3]).removesuffix(".0").removesuffix(".00"),
        "slot": "; ".join(sections), "instructor": "; ".join(instr), "co_instructors": "; ".join(co),
        "room": "; ".join(rooms), "students_registered": clean(rs[0][8]).removesuffix(".0"), "note": ""})
for alias, real in ALIASES.items():
    row = next(c for c in courses if c["course_code"] == real)
    courses.append({**row, "course_code": alias,
                    "note": f"Alias of {real} added for the assignment's example prompts; the official timetable lists this course as {real}."})
courses.sort(key=lambda c: c["course_code"])

# ---------- faculty.csv ----------
known = {}  # normalised name -> (name, department, email, source)
for r in csv.DictReader(open(SRC / "faculty_emails_collected.csv", encoding="utf-8-sig")):
    known[norm_name(r["name"])] = (r["name"].strip(), r["department"].strip(), r["email"].strip(),
                                   "faculty directory (collected)")
for r in wb["Time Table Committee Members"].iter_rows(min_row=3, values_only=True):
    # one row can hold two people: "Dr. Pankaj Chavan / Dr. Rajlaxmi Chouhan" | email1 | email2
    for name, email in zip(clean(r[2]).split("<")[0].split("/"), [r[3], r[4]]):
        name = re.sub(r"^Dr\.?\s*", "", name.strip()).strip()
        if name and email:
            known.setdefault(norm_name(name), (name, clean(r[1]), clean(email), "Time Table Committee sheet"))

person_depts = defaultdict(set)
for c in courses:
    for p in [x for x in c["instructor"].split("; ") + c["co_instructors"].split("; ") if x]:
        person_depts[norm_name(p)].add(c["department"])


def match_known(name):
    """Exact normalised match, else the one known name whose tokens are a subset/superset of this name's.
    Accepted only if the departments share a word, or the directory name is fully contained in the timetable
    name with 2+ tokens ('Tanmay Inamdar' ~ 'Tanmay Nitin Inamdar'); this rejects e.g. timetable 'Atul Kumar'
    (Chemistry) ~ directory 'Atul Kumar Sharma' (Mechanical)."""
    k = norm_name(name)
    if k in known:
        return k
    t = set(k.split())
    dwords = set(re.findall(r"[a-z]{4,}", " ".join(person_depts[k]).lower())) - {"school", "engineering"}
    cands = [kk for kk in known if (set(kk.split()) <= t or t <= set(kk.split())) and (
        dwords & set(re.findall(r"[a-z]{4,}", known[kk][1].lower()))
        or (set(kk.split()) <= t and len(set(kk.split())) >= 2))]
    return cands[0] if len(cands) == 1 else None


faculty, used = {}, set()
for c in courses:
    for p in [x for x in c["instructor"].split("; ") + c["co_instructors"].split("; ") if x]:
        k = norm_name(p)
        if k in faculty:
            continue
        kk = match_known(p)
        if kk:
            if kk != k:
                print(f"  matched timetable name {p!r} to directory name {known[kk][0]!r}")
            used.add(kk)
            _, dept, email, src = known[kk]
        else:
            dept = c["department"]
            email, src = ".".join(k.split()) + "@dummy.iitj.ac.in", "dummy placeholder (real address not collected)"
        faculty[k] = {"name": p, "department": dept or c["department"], "email": email, "email_source": src}
for k, (name, dept, email, src) in known.items():
    if k not in used and email not in {f["email"] for f in faculty.values()}:
        faculty.setdefault(k, {"name": name, "department": dept, "email": email, "email_source": src})


def write(path, rows):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    print(f"{path.relative_to(DATA.parent)}: {len(rows)} rows")


write(DATA / "slots.csv", slots)
write(DATA / "courses.csv", courses)
write(DATA / "faculty.csv", sorted(faculty.values(), key=lambda f: f["name"].lower()))

# ---------- docs/raw markdown version of the non-tabular sheets, so the RAG index can cite them ----------
md = ["# Slot System for B.Tech. 2nd year onwards, AY 2026-27 Semester I (official timetable workbook)", "",
      "## Slot timings", ""]
per_slot = defaultdict(list)
for s in slots:
    per_slot[s["slot"]].append(f"{s['day']} {s['start_time']}-{s['end_time']}")
md += [f"- Slot {k}: " + ", ".join(v) for k, v in per_slot.items()]
for sheet in ("Slot System", "Slots 2nd and 4th year common c", "Extra Slots for MTech Executive", "Examination",
              "Time Table Committee Members", "List of MinorSpecialization Coo"):
    md += ["", f"## {sheet.replace('Coo', 'Coordinators').replace('common c', 'common courses')}", ""]
    for r in wb[sheet].iter_rows(values_only=True):
        cells = [clean(v) for v in r if clean(v)]
        if cells:
            md.append(" | ".join(cells))
DOC.write_text("\n".join(md) + "\n")
print(f"{DOC.relative_to(DATA.parent)} written")
